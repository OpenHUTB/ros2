#!/usr/bin/env python3
"""Export the five Markdown documents to a Word file for format checking."""

import os
import sys
from pathlib import Path

import markdown
from docx import Document
from htmldocx import HtmlToDocx


PACKAGE_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[4]
DOCS_ROOT = (
    Path(os.environ['AUTO_PARKING_DOCS_DIR'])
    if os.environ.get('AUTO_PARKING_DOCS_DIR')
    else REPO_ROOT / 'docs'
)
SOURCE_DIR = DOCS_ROOT / 'ground' / 'mujoco_auto_parking'
OUTPUT_DIR = DOCS_ROOT / 'evidence' / 'mujoco_auto_parking'
FILES = [
    'index.md',
    '01_physical_simulation_keyboard.md',
    '02_perception_trajectory_control.md',
    '03_slam_navigation.md',
    '04_end_to_end.md',
]


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading('MuJoCo 四轮差速车自动泊车', level=0)
    document.add_paragraph(
        '说明：LaTeX 公式在 Word 中保留为文本；版式校验以标题层级、表格和代码块为准。'
    )
    parser = HtmlToDocx()
    for name in FILES:
        path = SOURCE_DIR / name
        if not path.is_file():
            raise FileNotFoundError(path)
        text = path.read_text(encoding='utf-8')
        html = markdown.markdown(
            text, extensions=['tables', 'fenced_code', 'sane_lists']
        )
        parser.add_html_to_document(html, document)
    output = OUTPUT_DIR / 'mujoco_auto_parking_report.docx'
    document.save(output)

    reopened = Document(output)
    print(f'Word 已导出: {output}')
    print(f'段落数: {len(reopened.paragraphs)}，表格数: {len(reopened.tables)}')
    assert len(reopened.paragraphs) > 100
    assert len(reopened.tables) > 5


if __name__ == '__main__':
    main()
